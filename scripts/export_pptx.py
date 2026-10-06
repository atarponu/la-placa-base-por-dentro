"""Genera exports/presentation.pptx con una diapositiva por escena (imagen a sangre).
Requiere: pip install python-pptx   ·   Ejecutar después de scripts/export.mjs
Nota: es una exportación estática; la versión principal es index.html (interactiva)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

root = Path(__file__).resolve().parent.parent
slides = sorted((root / 'exports/slides').glob('*.jpg'))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
for img in slides:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
prs.save(root / 'exports/presentation.pptx')
print('PPTX →', root / 'exports/presentation.pptx', len(slides), 'diapositivas')
